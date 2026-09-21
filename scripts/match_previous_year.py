#!/usr/bin/env python3
"""Match last-year Excel follow-up reports onto current SCHOOL_DATA students.

Year mapping (last year → this year):
  Pre-schoolers → Primary Year 1
  1ST Graders   → Primary Year 2
  2ND Graders   → Primary Year 3
  3RD Graders   → Primary Year 4
  4TH Graders   → Primary Year 5
  5TH Graders   → Middle Year 1

Writes:
  data/school_data.json
  js/school-data.js
  data/previous_year_match_report.json
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError as e:
    raise SystemExit("Install openpyxl first: pip install openpyxl") from e

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "school_data.json"
JS_PATH = ROOT / "js" / "school-data.js"
REPORT_PATH = ROOT / "data" / "previous_year_match_report.json"

HEADER_MAP = {
    "الرقم": "number",
    "اسم الطالب": "fullName",
    "تاريخ الميلاد": "dateOfBirth",
    "العمر": "age",
    "المعدل السنوي": "annualAverage",
    "الحالة الإجتماعية": "socialStatus",
    "الحالة الاجتماعية": "socialStatus",
    "الحالة الصحية": "healthStatus",
    "نقاط القوة": "strengths",
    "نقاط الضعف": "weaknesses",
    "الأداء السلوكي": "behavioralPerformance",
    "الأداء التحصيلي": "academicPerformance",
    "الأداء الدراسي": "academicPerformance",
    "المواهب (إن وجدت)": "talents",
    "ملاحظات إضافية": "additionalNotes",
    "خطة العمل": "actionPlan",
    "ملاحظة عامة": "generalNote",
}

FILE_MAP = [
    ("Pre-schoolers Report.xlsx", "primary", 1, "Pre-school", "السنة التحضيرية"),
    ("1ST Graders.xlsx", "primary", 2, "Year 1", "السنة الأولى"),
    ("2ND Graders.xlsx", "primary", 3, "Year 2", "السنة الثانية"),
    ("3RD Graders.xlsx", "primary", 4, "Year 3", "السنة الثالثة"),
    ("4TH Graders.xlsx", "primary", 5, "Year 4", "السنة الرابعة"),
    ("5TH Graders.xlsx", "middle", 1, "Year 5 (Primary)", "السنة الخامسة ابتدائي"),
]

ALEF = re.compile(r"[أإآٱ]")
YEH = re.compile(r"[ىي]")
TEH_MARBUTA = re.compile(r"ة")
TATWEEL = re.compile(r"ـ")
DIACRITICS = re.compile(r"[\u064B-\u065F\u0670]")
NON_LETTER = re.compile(r"[^\w\u0600-\u06FF]+", re.UNICODE)
DETAIL_KEYS = [
    "number",
    "fullName",
    "dateOfBirth",
    "age",
    "annualAverage",
    "socialStatus",
    "healthStatus",
    "strengths",
    "weaknesses",
    "behavioralPerformance",
    "academicPerformance",
    "talents",
    "additionalNotes",
    "actionPlan",
    "generalNote",
]


def normalize_name(s: str) -> str:
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", str(s).strip())
    s = DIACRITICS.sub("", s)
    s = TATWEEL.sub("", s)
    s = ALEF.sub("ا", s)
    s = YEH.sub("ي", s)
    s = TEH_MARBUTA.sub("ه", s)
    s = s.replace("ؤ", "و").replace("ئ", "ي")
    s = s.replace("عبدال", "عبد ال").replace("بنال", "بن ال")
    s = NON_LETTER.sub(" ", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def name_tokens(s: str) -> list[str]:
    return [t for t in normalize_name(s).split() if t]


def skip_particle(t: str) -> bool:
    return t in {"بن", "ابن", "بنت", "ال", "عبد"}


def significant_tokens(toks: list[str]) -> list[str]:
    return [t for t in toks if not skip_particle(t) and len(t) > 1]


def fmt_dob(v) -> str:
    if v is None or v == "" or v == "/":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if not s or s == "/":
        return ""
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            pass
    return s


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return fmt_dob(v)
    s = str(v).strip()
    if s in ("/", "None", "nan"):
        return ""
    return s


def parse_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    students = []
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        if not rows:
            continue
        header_idx = None
        headers = None
        for i, row in enumerate(rows[:5]):
            vals = [cell_str(c) for c in (row or [])]
            if any(v == "اسم الطالب" for v in vals):
                header_idx = i
                headers = vals
                break
        if header_idx is None:
            continue
        colmap = {}
        for ci, h in enumerate(headers):
            key = HEADER_MAP.get(h)
            if key:
                colmap[ci] = key
        name_ci = next((ci for ci, k in colmap.items() if k == "fullName"), 1)
        for row in rows[header_idx + 1 :]:
            if not row:
                continue
            name = cell_str(row[name_ci] if name_ci < len(row) else None)
            if not name:
                continue
            rec = {k: "" for k in DETAIL_KEYS}
            for ci, key in colmap.items():
                if ci >= len(row):
                    continue
                val = row[ci]
                rec[key] = fmt_dob(val) if key == "dateOfBirth" else cell_str(val)
            rec["fullName"] = name
            rec["priorClass"] = sn
            students.append(rec)
    wb.close()
    return students


def lev(a: str, b: str) -> int:
    if a == b:
        return 0
    la, lb = len(a), len(b)
    if abs(la - lb) > 2:
        return 99
    if la < lb:
        a, b, la, lb = b, a, lb, la
    prev = list(range(lb + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (0 if ca == cb else 1)))
        prev = cur
    return prev[lb]


def tokens_soft_equal(a: str, b: str) -> bool:
    if a == b:
        return True
    if not a or not b:
        return False
    if a.startswith(b) or b.startswith(a):
        return min(len(a), len(b)) >= 3
    d = lev(a, b)
    mx = max(len(a), len(b))
    if mx <= 7:
        return d <= 1
    return d <= 2


def match_score(cur_toks: list[str], prior_toks: list[str]) -> int:
    if not cur_toks or not prior_toks:
        return 0
    if cur_toks == prior_toks:
        return 100
    if set(cur_toks) == set(prior_toks):
        return 98
    cs = significant_tokens(cur_toks)
    ps = significant_tokens(prior_toks)
    if not cs or not ps:
        return 0
    used = set()
    matched = 0
    for c in cs:
        best_j, best_d = None, 99
        for j, p in enumerate(ps):
            if j in used:
                continue
            if tokens_soft_equal(c, p):
                d = 0 if c == p else lev(c, p)
                if d < best_d:
                    best_d, best_j = d, j
        if best_j is not None:
            used.add(best_j)
            matched += 1
    cov_short = matched / min(len(cs), len(ps))
    cov_long = matched / max(len(cs), len(ps))
    if cov_short < 0.8:
        return 0
    if matched < 2 and max(len(cs), len(ps)) >= 2:
        return 0
    return min(99, int(50 + 35 * cov_short + 15 * cov_long))


def level_year(lv: dict) -> int | None:
    y = lv.get("year")
    if y is not None:
        try:
            return int(y)
        except (TypeError, ValueError):
            pass
    m = re.search(r"(\d+)", str(lv.get("id", "")) + " " + str(lv.get("name", "")))
    return int(m.group(1)) if m else None


def clear_enrichment(data: dict) -> None:
    for dept_id in ("primary", "middle"):
        for lv in data[dept_id]["levels"]:
            for c in lv["classes"]:
                for s in c["students"]:
                    s.pop("previousYearDetails", None)
                    s["dateOfBirth"] = ""


def run(excel_dir: Path) -> dict:
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    clear_enrichment(data)

    prior_by_target = {}
    for fname, dept, year, prev_en, prev_ar in FILE_MAP:
        path = excel_dir / fname
        if not path.exists():
            raise FileNotFoundError(path)
        studs = parse_workbook(path)
        prior_by_target[(dept, year)] = {
            "previousYearLabelEn": prev_en,
            "previousYearLabelAr": prev_ar,
            "sourceFile": fname,
            "students": studs,
        }
        print(f"{fname}: {len(studs)} → {dept} Y{year}")

    stats = {
        "exact": 0,
        "fuzzy": 0,
        "unmatched_current": 0,
        "unmatched_prior": 0,
        "ambiguous": 0,
    }
    match_report = []

    for (dept_id, year), bundle in prior_by_target.items():
        prior_students = bundle["students"]
        by_exact = defaultdict(list)
        prior_meta = []
        for p in prior_students:
            n = normalize_name(p["fullName"])
            toks = name_tokens(p["fullName"])
            by_exact[n].append(p)
            prior_meta.append((n, toks, p))

        used = set()
        current = []
        for lv in data[dept_id]["levels"]:
            if level_year(lv) != year:
                continue
            for c in lv["classes"]:
                current.extend(c["students"])

        for s in current:
            cn = normalize_name(s["fullName"])
            cands = [p for p in by_exact.get(cn, []) if id(p) not in used]
            if len(cands) == 1:
                s["_match"] = (cands[0], "exact")
                used.add(id(cands[0]))
                stats["exact"] += 1
            elif len(cands) > 1:
                s["_match"] = (cands[0], "exact-amb")
                used.add(id(cands[0]))
                stats["exact"] += 1
                stats["ambiguous"] += 1

        for s in current:
            if s.get("_match"):
                continue
            ct = name_tokens(s["fullName"])
            best = []
            for _n, pt, p in prior_meta:
                if id(p) in used:
                    continue
                sc = match_score(ct, pt)
                if sc >= 82:
                    best.append((sc, p))
            best.sort(key=lambda x: -x[0])
            if not best:
                continue
            if len(best) == 1 or best[0][0] > best[1][0] + 2:
                s["_match"] = (best[0][1], f"fuzzy:{best[0][0]}")
                used.add(id(best[0][1]))
                stats["fuzzy"] += 1
            else:
                stats["ambiguous"] += 1

        for s in current:
            m = s.pop("_match", None)
            if not m:
                stats["unmatched_current"] += 1
                match_report.append(
                    {
                        "current": s["fullName"],
                        "prior": None,
                        "dept": dept_id,
                        "year": year,
                        "method": "unmatched",
                    }
                )
                continue
            match, method = m
            details = {
                "previousYear": bundle["previousYearLabelEn"],
                "previousYearAr": bundle["previousYearLabelAr"],
                "previousClass": match.get("priorClass", ""),
                "sourceFile": bundle["sourceFile"],
                "matchedName": match["fullName"],
                "matchMethod": method,
                "number": match.get("number", ""),
                "dateOfBirth": match.get("dateOfBirth", ""),
                "age": match.get("age", ""),
                "annualAverage": match.get("annualAverage", ""),
                "socialStatus": match.get("socialStatus", ""),
                "healthStatus": match.get("healthStatus", ""),
                "strengths": match.get("strengths", ""),
                "weaknesses": match.get("weaknesses", ""),
                "behavioralPerformance": match.get("behavioralPerformance", ""),
                "academicPerformance": match.get("academicPerformance", ""),
                "talents": match.get("talents", ""),
                "additionalNotes": match.get("additionalNotes", ""),
                "actionPlan": match.get("actionPlan", ""),
                "generalNote": match.get("generalNote", ""),
            }
            s["previousYearDetails"] = details
            if match.get("dateOfBirth"):
                s["dateOfBirth"] = match["dateOfBirth"]
            match_report.append(
                {
                    "current": s["fullName"],
                    "prior": match["fullName"],
                    "dept": dept_id,
                    "year": year,
                    "method": method,
                    "dob": match.get("dateOfBirth", ""),
                }
            )

        unused = [p for p in prior_students if id(p) not in used]
        stats["unmatched_prior"] += len(unused)
        matched_n = sum(1 for s in current if s.get("previousYearDetails"))
        print(f"{dept_id} Y{year}: matched {matched_n}/{len(current)} (unused prior {len(unused)})")

    DATA_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    JS_PATH.write_text(
        "/** Auto-generated from data/school_data.json — do not edit by hand */\n"
        "const SCHOOL_DATA = "
        + json.dumps(data, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    REPORT_PATH.write_text(
        json.dumps({"stats": stats, "matches": match_report}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print("STATS", stats)
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "excel_dir",
        type=Path,
        help="Folder containing Pre-schoolers / 1ST–5TH Graders.xlsx",
    )
    args = parser.parse_args()
    run(args.excel_dir.resolve())


if __name__ == "__main__":
    main()
