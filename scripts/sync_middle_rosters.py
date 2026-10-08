#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Sync college (middle) rosters from قائمة المتوسط.xlsx into the live site.

Excel is authoritative for class membership, Arabic names, DOB, gender, number.
Preserves student IDs when matched. Soft-flags site extras (no delete unless --delete-extra).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from collections import Counter
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".pydeps"))
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402
from bilingual_names import latin_parts_from_ar  # noqa: E402

API = "https://quality-education-algerie.duckdns.org/api"
ADMIN_PASSWORD = "AdminBBC2026"
APPLY = "--apply" in sys.argv
DELETE_EXTRA = "--delete-extra" in sys.argv
DOWNLOADS = Path("/Users/yahyaabderrahmanemahdi/Downloads")
REPORT = ROOT / "reports" / "middle_roster_sync_report.json"

PARTICLES = {
    "بن",
    "ابن",
    "بنت",
    "ال",
    "عبد",
    "بو",
    "ايت",
    "آيت",
    "أيت",
    "حاج",
    "ولد",
    "سيدي",
    "ابو",
    "أبو",
}


def find_excel() -> Path:
    for p in DOWNLOADS.iterdir():
        if p.suffix.lower() == ".xlsx" and "المتوسط" in p.name:
            return p
    raise FileNotFoundError("قائمة المتوسط.xlsx not found in Downloads")


def req(path, data=None, token=None, method=None):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(data).encode("utf-8") if data is not None else None
    m = method or ("POST" if body is not None else "GET")
    request = urllib.request.Request(API + path, data=body, headers=headers, method=m)
    try:
        with urllib.request.urlopen(request, timeout=180) as res:
            return json.loads(res.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        err = e.read().decode("utf-8", "replace")
        raise RuntimeError(f"HTTP {e.code}: {err}") from e


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", str(s or "").strip())
    s = re.sub(r"[\u064B-\u065F\u0670]", "", s)
    s = s.replace("ـ", "")
    for a, b in {
        "أ": "ا",
        "إ": "ا",
        "آ": "ا",
        "ٱ": "ا",
        "ى": "ي",
        "ة": "ه",
        "ؤ": "و",
        "ئ": "ي",
        "ء": "",
    }.items():
        s = s.replace(a, b)
    s = s.replace("عبدال", "عبد ال").replace("بنال", "بن ال")
    s = re.sub(r"[^\w\u0600-\u06FF]+", " ", s, flags=re.UNICODE)
    return re.sub(r"\s+", " ", s).strip().lower()


def toks(s: str) -> list[str]:
    return [t for t in norm(s).split() if t]


def sig(tok_list: list[str]) -> list[str]:
    return [t for t in tok_list if t not in PARTICLES and len(t) > 1]


def gender_map(g) -> str:
    g = str(g or "").strip()
    if "أنث" in g or "انث" in g:
        return "Female"
    if "ذكر" in g:
        return "Male"
    return ""


def dob_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    return s


def class_id_for(year: int, class_num: int) -> str:
    return f"M-Y{year}-C{class_num:02d}"


def site_code_for(year: int, class_num: int) -> str:
    return f"{year}M{class_num}"


def parse_sheet_name(name: str) -> tuple[int, int] | None:
    m = re.match(r"^\s*(\d)\s*م\s*(\d+)\s*$", str(name or "").strip())
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def parse_excel(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    out = []
    for sn in wb.sheetnames:
        parsed = parse_sheet_name(sn)
        if not parsed:
            continue
        year, class_num = parsed
        ws = wb[sn]
        header_row = None
        for r in range(1, 20):
            vals = [str(ws.cell(r, c).value or "") for c in range(1, 8)]
            if any("لقب" in v for v in vals):
                header_row = r
                break
        if not header_row:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            num = ws.cell(r, 1).value
            last = ws.cell(r, 2).value
            first = ws.cell(r, 3).value
            dob = ws.cell(r, 4).value
            gender = ws.cell(r, 5).value
            notes = ws.cell(r, 6).value
            last_s = str(last or "").strip()
            first_s = str(first or "").strip()
            if not last_s and not first_s:
                continue
            # Stop at class-teacher footer blocks
            blob = f"{last_s} {first_s}"
            if any(
                x in blob
                for x in (
                    "قائمة اساتذ",
                    "قائمة أساتذ",
                    "ل,عربية",
                    "ل،عربية",
                    "اجتماعيات",
                    "الاساتذ",
                    "الأساتذ",
                )
            ):
                break
            if isinstance(num, str) and not str(num).strip().isdigit():
                continue
            raw = f"{last_s} {first_s}".strip()
            try:
                number = int(num)
            except Exception:
                number = None
            if number is None:
                # numbered list ended; ignore trailing teacher labels
                continue
            out.append(
                {
                    "year": year,
                    "classNum": class_num,
                    "classId": class_id_for(year, class_num),
                    "classCode": site_code_for(year, class_num),
                    "number": number,
                    "rawFull": raw,
                    "lastName": last_s,
                    "firstName": first_s,
                    "fullName": raw,
                    "gender": gender_map(gender),
                    "dateOfBirth": dob_str(dob),
                    "notes": str(notes or "").strip(),
                    "sourceFile": path.name,
                }
            )
    return out


def score(ex: dict, st: dict) -> tuple[float, str]:
    en, sn_ = norm(ex["rawFull"]), norm(st["fullName"])
    if not en or not sn_:
        return 0, "empty"
    dob_ex = (ex.get("dateOfBirth") or "")[:10]
    dob_st = str(st.get("dateOfBirth") or "")[:10]
    dob_bonus = 0.08 if dob_ex and dob_st and dob_ex == dob_st else 0.0

    if en == sn_:
        return min(1.0, 0.95 + dob_bonus), "exact"
    es, ss = set(sig(toks(ex["rawFull"]))), set(sig(toks(st["fullName"])))
    if es and es == ss:
        return min(1.0, 0.92 + dob_bonus), "sig-equal"
    if es and ss and (es <= ss or ss <= es) and len(es & ss) >= 2:
        return min(1.0, 0.88 + dob_bonus), "sig-subset"
    if es and ss and len(es & ss) >= 2:
        j = len(es & ss) / len(es | ss)
        if j >= 0.75:
            return min(1.0, 0.82 + j * 0.1 + dob_bonus), "sig-jaccard"
    # same DOB + strong fuzzy name
    ratio = SequenceMatcher(None, en, sn_).ratio()
    compact = SequenceMatcher(None, en.replace(" ", ""), sn_.replace(" ", "")).ratio()
    ratio = max(ratio, compact)
    if dob_bonus and ratio >= 0.55:
        return min(1.0, 0.78 + ratio * 0.15 + dob_bonus), "dob+fuzzy"
    if ratio >= 0.88:
        return min(1.0, ratio + dob_bonus * 0.5), "fuzzy"
    if ex.get("classId") == st.get("classId") and ratio >= 0.75:
        return min(1.0, 0.72 + ratio * 0.2), "same-class-fuzzy"
    return ratio * 0.5, "weak"


def build_plan(excel: list[dict], site: list[dict]):
    candidates = []
    for i, ex in enumerate(excel):
        for st in site:
            sc, method = score(ex, st)
            if sc < 0.72:
                continue
            candidates.append((sc, method, i, st["id"]))
    candidates.sort(reverse=True)

    used_ex, used_st, pairs = set(), set(), []
    for sc, method, i, sid in candidates:
        if i in used_ex or sid in used_st:
            continue
        used_ex.add(i)
        used_st.add(sid)
        pairs.append((excel[i], next(s for s in site if s["id"] == sid), sc, method))

    # recovery: same year high fuzzy
    for i in [j for j in range(len(excel)) if j not in used_ex]:
        ex = excel[i]
        best, best_sc, best_m = None, 0.0, ""
        for st in site:
            if st["id"] in used_st:
                continue
            if st["year"] != ex["year"]:
                continue
            sc, method = score(ex, st)
            if sc > best_sc:
                best, best_sc, best_m = st, sc, method
        if best and best_sc >= 0.80:
            used_ex.add(i)
            used_st.add(best["id"])
            pairs.append((ex, best, best_sc, "recover-" + best_m))

    updates, inserts = [], []
    for ex, st, sc, method in pairs:
        last = (ex.get("lastName") or "").strip()
        first = (ex.get("firstName") or "").strip()
        full = f"{last} {first}".strip()
        name_changed = (
            full != (st.get("fullName") or "")
            or last != (st.get("lastName") or "")
            or first != (st.get("firstName") or "")
        )
        class_changed = ex["classId"] != st["classId"]
        gender_changed = bool(ex["gender"]) and ex["gender"] != (st.get("gender") or "")
        number_changed = ex["number"] is not None and ex["number"] != st.get("number")
        dob_ex = ex.get("dateOfBirth") or ""
        dob_st = str(st.get("dateOfBirth") or "")[:10]
        dob_changed = bool(dob_ex) and dob_ex != dob_st

        if name_changed:
            fl, ll, ful, _ = latin_parts_from_ar(first, last, full)
        else:
            fl = st.get("firstNameLatin") or ""
            ll = st.get("lastNameLatin") or ""
            ful = st.get("fullNameLatin") or ""
            if not ful:
                fl, ll, ful, _ = latin_parts_from_ar(first, last, full)

        if name_changed or class_changed or gender_changed or number_changed or dob_changed:
            updates.append(
                {
                    "studentId": st["id"],
                    "fromName": st["fullName"],
                    "fullName": full,
                    "lastName": last,
                    "firstName": first,
                    "firstNameLatin": fl,
                    "lastNameLatin": ll,
                    "fullNameLatin": ful,
                    "searchName": f"{full} {ful} {first} {fl} {last} {ll}".strip(),
                    "fromClass": st["classId"],
                    "classId": ex["classId"],
                    "departmentId": "middle",
                    "number": ex["number"] if ex["number"] is not None else st.get("number"),
                    "gender": ex["gender"] or st.get("gender") or "",
                    "dateOfBirth": dob_ex or dob_st or "",
                    "nameChanged": name_changed,
                    "classChanged": class_changed,
                    "score": round(float(sc), 3),
                    "method": method,
                    "excelRaw": ex["rawFull"],
                }
            )

    for i, ex in enumerate(excel):
        if i in used_ex:
            continue
        last, first, full = ex["lastName"], ex["firstName"], ex["fullName"]
        fl, ll, ful, _ = latin_parts_from_ar(first, last, full)
        inserts.append(
            {
                **ex,
                "firstNameLatin": fl,
                "lastNameLatin": ll,
                "fullNameLatin": ful,
                "departmentId": "middle",
                "searchName": f"{full} {ful} {first} {fl} {last} {ll}".strip(),
            }
        )

    site_left = [s for s in site if s["id"] not in used_st]
    return pairs, updates, inserts, site_left


def main():
    path = find_excel()
    print(f"Parsing {path.name}…")
    excel = parse_excel(path)
    print(f"  excel students: {len(excel)} {dict(Counter(e['year'] for e in excel))}")

    print("Fetching live school-data…")
    login = req("/auth/login", {"password": ADMIN_PASSWORD, "role": "admin"})
    token = login["token"]
    school = req("/school-data", token=token)
    site = []
    for lev in school["middle"]["levels"]:
        for cls in lev["classes"]:
            for s in cls["students"]:
                site.append(
                    {
                        **s,
                        "year": lev["id"],
                        "classCode": cls.get("code"),
                        "classId": cls["id"],
                    }
                )
    print(f"  live middle: {len(site)}")

    pairs, updates, inserts, site_left = build_plan(excel, site)
    name_updates = [u for u in updates if u["nameChanged"]]
    class_moves = [u for u in updates if u["classChanged"]]

    print(
        f"pairs={len(pairs)} updates={len(updates)} "
        f"name_corrections={len(name_updates)} class_moves={len(class_moves)} "
        f"inserts={len(inserts)} site_unmatched={len(site_left)}"
    )
    print("\nName corrections:")
    for u in name_updates[:40]:
        print(
            f"  [{u['score']}/{u['method']}] {u['fromClass']}→{u['classId']} "
            f"{u['fromName']!r} → {u['fullName']!r}"
        )
    if len(name_updates) > 40:
        print(f"  … +{len(name_updates) - 40} more")
    print(f"\nInserts ({len(inserts)}):")
    for ex in inserts:
        print(f"  Y{ex['year']} {ex['classCode']} #{ex['number']} {ex['fullName']!r}")
    print(f"\nLeft on site (not in Excel): {len(site_left)}")
    for s in site_left:
        print(f"  Y{s['year']} {s['classCode']} {s['fullName']} ({s['id']})")

    report = {
        "stats": {
            "excel": len(excel),
            "live": len(site),
            "pairs": len(pairs),
            "updates": len(updates),
            "name_corrections": len(name_updates),
            "class_moves": len(class_moves),
            "inserts": len(inserts),
            "site_unmatched": len(site_left),
            "applied": False,
            "deleted_extra": 0,
        },
        "name_corrections": name_updates,
        "updates": updates,
        "inserts": inserts,
        "site_unmatched": [
            {
                "id": s["id"],
                "year": s["year"],
                "class": s["classCode"],
                "fullName": s["fullName"],
            }
            for s in site_left
        ],
    }

    if APPLY:
        print("\nApplying to live API…")
        ok_u = ok_i = err = skip_i = ok_d = 0
        existing_keys = set()
        for lev in school["middle"]["levels"]:
            for cls in lev["classes"]:
                for s in cls["students"]:
                    existing_keys.add((norm(s.get("fullName") or ""), cls["id"], s.get("number")))
        for u in updates:
            body = {
                "firstName": u["firstName"],
                "lastName": u["lastName"],
                "fullName": u["fullName"],
                "firstNameLatin": u["firstNameLatin"],
                "lastNameLatin": u["lastNameLatin"],
                "fullNameLatin": u["fullNameLatin"],
                "searchName": u["searchName"],
                "classId": u["classId"],
                "departmentId": "middle",
                "number": u["number"],
                "gender": u["gender"],
                "dateOfBirth": u.get("dateOfBirth") or "",
            }
            try:
                req(f"/students/{u['studentId']}", body, token=token, method="PUT")
                ok_u += 1
            except Exception as e:
                err += 1
                print(f"  UPDATE FAIL {u['studentId']} {u['fromName']}: {e}")
        for ex in inserts:
            key = (norm(ex["fullName"]), ex["classId"], ex["number"])
            if key in existing_keys:
                skip_i += 1
                continue
            body = {
                "firstName": ex["firstName"],
                "lastName": ex["lastName"],
                "fullName": ex["fullName"],
                "firstNameLatin": ex["firstNameLatin"],
                "lastNameLatin": ex["lastNameLatin"],
                "fullNameLatin": ex["fullNameLatin"],
                "searchName": ex["searchName"],
                "classId": ex["classId"],
                "departmentId": "middle",
                "number": ex["number"] or 0,
                "gender": ex["gender"],
                "dateOfBirth": ex.get("dateOfBirth") or "",
                "notes": ex.get("notes") or "",
            }
            try:
                created = req("/students", body, token=token, method="POST")
                ok_i += 1
                existing_keys.add(key)
                if isinstance(created, dict) and created.get("id"):
                    existing_keys.add(
                        (
                            norm(created.get("fullName") or ex["fullName"]),
                            ex["classId"],
                            ex["number"],
                        )
                    )
            except Exception as e:
                err += 1
                print(f"  INSERT FAIL {ex['fullName']}: {e}")

        deleted_rows = []
        if DELETE_EXTRA and site_left:
            print(f"\nDeleting {len(site_left)} students not on Excel…")
            for s in site_left:
                try:
                    req(f"/students/{s['id']}", token=token, method="DELETE")
                    ok_d += 1
                    deleted_rows.append(
                        {"id": s["id"], "class": s["classCode"], "fullName": s["fullName"]}
                    )
                    print(f"  DEL {s['classCode']} {s['fullName']}")
                except Exception as e:
                    err += 1
                    print(f"  DEL FAIL {s['id']} {s['fullName']}: {e}")

        print(
            f"Applied updates={ok_u} inserts={ok_i} deleted={ok_d} "
            f"skipped_inserts={skip_i} errors={err}"
        )
        report["stats"]["applied"] = True
        report["stats"]["applied_updates"] = ok_u
        report["stats"]["applied_inserts"] = ok_i
        report["stats"]["deleted_extra"] = ok_d
        report["stats"]["skipped_inserts"] = skip_i
        report["stats"]["errors"] = err
        report["deleted"] = deleted_rows

        print("Refreshing local school_data from live…")
        school2 = req("/school-data", token=token)
        data_path = ROOT / "data" / "school_data.json"
        data_path.write_text(json.dumps(school2, ensure_ascii=False, indent=2), encoding="utf-8")
        js_path = ROOT / "js" / "school-data.js"
        js_path.write_text(
            "window.SCHOOL_DATA = " + json.dumps(school2, ensure_ascii=False) + ";\n",
            encoding="utf-8",
        )
        print(f"  wrote {data_path} and {js_path}")

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nReport: {REPORT}")
    if not APPLY:
        print("Dry-run only. Re-run with --apply to write to the live site.")


if __name__ == "__main__":
    main()
