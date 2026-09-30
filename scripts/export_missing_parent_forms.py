#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build an organized Excel of students whose parents have not submitted the form yet.
Ready to upload to Google Sheets."""
from __future__ import annotations

import json
import urllib.request
import urllib.error
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

API = "https://quality-education-algerie.duckdns.org/api"
ADMIN_PASSWORD = "AdminBBC2026"
OUT = Path(__file__).resolve().parents[1] / "reports" / "missing_parent_form_submissions.xlsx"


def req(path, data=None, token=None):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(data).encode("utf-8") if data is not None else None
    method = "POST" if body is not None else "GET"
    request = urllib.request.Request(API + path, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=180) as res:
            return json.loads(res.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        err = e.read().decode("utf-8", "replace")
        raise RuntimeError(f"HTTP {e.code}: {err}") from e


def style_header(ws, n_cols):
    fill = PatternFill("solid", fgColor="0B1F3A")
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, n_cols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ws.max_row}"


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    login = req("/auth/login", {"password": ADMIN_PASSWORD, "role": "admin"})
    token = login["token"]
    school = req("/school-data", token=token)

    rows = []
    linked = 0
    total = 0
    for dept_key, dept_label in (("primary", "Primary"), ("middle", "Middle / College")):
        dept = school.get(dept_key) or {}
        for level in dept.get("levels") or []:
            level_name = level.get("name") or level.get("id") or ""
            year = level.get("year") or ""
            for cls in level.get("classes") or []:
                code = cls.get("code") or ""
                for s in cls.get("students") or []:
                    total += 1
                    if s.get("parentFormId") or s.get("parentProfile"):
                        linked += 1
                        continue
                    rows.append(
                        {
                            "department": dept_label,
                            "deptKey": dept_key,
                            "year": year or cls.get("year") or "",
                            "level": level_name,
                            "classCode": code,
                            "roster": s.get("number") if s.get("number") is not None else "",
                            "lastNameAr": s.get("lastName") or "",
                            "firstNameAr": s.get("firstName") or "",
                            "lastNameLatin": s.get("lastNameLatin") or "",
                            "firstNameLatin": s.get("firstNameLatin") or "",
                            "dob": s.get("dateOfBirth") or "",
                            "gender": s.get("gender") or "",
                            "studentId": s.get("id") or "",
                        }
                    )

    rows.sort(
        key=lambda r: (
            0 if r["deptKey"] == "primary" else 1,
            str(r["year"]),
            r["classCode"],
            r["lastNameLatin"] or r["lastNameAr"],
            r["firstNameLatin"] or r["firstNameAr"],
        )
    )

    wb = Workbook()

    # Summary
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Missing parent form submissions"
    summary["A1"].font = Font(bold=True, size=14, color="0B1F3A")
    summary["A3"] = "Total students"
    summary["B3"] = total
    summary["A4"] = "Parent form linked"
    summary["B4"] = linked
    summary["A5"] = "Missing submissions"
    summary["B5"] = len(rows)
    summary["A6"] = "Coverage %"
    summary["B6"] = round(100 * linked / total, 1) if total else 0

    by_dept = defaultdict(int)
    by_year = defaultdict(int)
    by_class = defaultdict(int)
    for r in rows:
        by_dept[r["department"]] += 1
        by_year[(r["department"], str(r["year"]), r["level"])] += 1
        by_class[(r["department"], r["classCode"])] += 1

    summary["A8"] = "Missing by department"
    summary["A8"].font = Font(bold=True)
    i = 9
    for k, v in sorted(by_dept.items()):
        summary[f"A{i}"] = k
        summary[f"B{i}"] = v
        i += 1
    i += 1
    summary[f"A{i}"] = "Missing by year / level"
    summary[f"A{i}"].font = Font(bold=True)
    i += 1
    summary[f"A{i}"] = "Department"
    summary[f"B{i}"] = "Year"
    summary[f"C{i}"] = "Level"
    summary[f"D{i}"] = "Missing"
    for c in range(1, 5):
        summary.cell(i, c).font = Font(bold=True)
    i += 1
    for (dept, year, level), v in sorted(by_year.items()):
        summary[f"A{i}"] = dept
        summary[f"B{i}"] = year
        summary[f"C{i}"] = level
        summary[f"D{i}"] = v
        i += 1
    autosize(summary, [28, 12, 28, 12])

    # All missing
    all_ws = wb.create_sheet("All missing")
    headers = [
        "#",
        "Department",
        "Year",
        "Level",
        "Class",
        "Roster #",
        "Last name (AR)",
        "First name (AR)",
        "Last name (Latin)",
        "First name (Latin)",
        "Date of birth",
        "Gender",
        "Student ID",
    ]
    all_ws.append(headers)
    for idx, r in enumerate(rows, 1):
        all_ws.append(
            [
                idx,
                r["department"],
                r["year"],
                r["level"],
                r["classCode"],
                r["roster"],
                r["lastNameAr"],
                r["firstNameAr"],
                r["lastNameLatin"],
                r["firstNameLatin"],
                r["dob"],
                r["gender"],
                r["studentId"],
            ]
        )
    style_header(all_ws, len(headers))
    autosize(all_ws, [6, 16, 8, 18, 10, 10, 18, 18, 18, 18, 14, 10, 16])

    # By department sheets
    for dept_key, dept_label, title in (
        ("primary", "Primary", "Primary missing"),
        ("middle", "Middle / College", "College missing"),
    ):
        ws = wb.create_sheet(title)
        ws.append(headers)
        n = 0
        for r in rows:
            if r["deptKey"] != dept_key:
                continue
            n += 1
            ws.append(
                [
                    n,
                    r["department"],
                    r["year"],
                    r["level"],
                    r["classCode"],
                    r["roster"],
                    r["lastNameAr"],
                    r["firstNameAr"],
                    r["lastNameLatin"],
                    r["firstNameLatin"],
                    r["dob"],
                    r["gender"],
                    r["studentId"],
                ]
            )
        style_header(ws, len(headers))
        autosize(ws, [6, 16, 8, 18, 10, 10, 18, 18, 18, 18, 14, 10, 16])

    # By class summary
    class_ws = wb.create_sheet("By class")
    class_ws.append(["Department", "Class", "Missing count"])
    for (dept, code), v in sorted(by_class.items(), key=lambda x: (x[0][0], x[0][1])):
        class_ws.append([dept, code, v])
    style_header(class_ws, 3)
    autosize(class_ws, [18, 12, 14])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"total={total} linked={linked} missing={len(rows)}")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
