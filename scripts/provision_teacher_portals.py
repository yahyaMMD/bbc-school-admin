#!/usr/bin/env python3
"""Provision teacher portal passwords and write an Excel credentials sheet."""
from __future__ import annotations

import json
import sys
import urllib.request
import urllib.error
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

API = "https://quality-education-algerie.duckdns.org/api"
OUT = Path(__file__).resolve().parents[1] / "reports" / "teacher_portal_credentials.xlsx"
ADMIN_PASSWORD = "AdminBBC2026"


def req(path: str, data=None, token: str | None = None):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(data).encode("utf-8") if data is not None else None
    method = "POST" if body is not None else "GET"
    request = urllib.request.Request(API + path, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=300) as res:
            return json.loads(res.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        err = e.read().decode("utf-8", "replace")
        raise RuntimeError(f"HTTP {e.code}: {err}") from e


def main():
    login = req("/auth/login", {"password": ADMIN_PASSWORD, "role": "admin"})
    token = login["token"]
    print("logged in as", login.get("role"))

    result = req("/teachers/provision-portals", {}, token=token)
    creds = result.get("credentials") or []
    print(f"provisioned {len(creds)} teachers")

    wb = Workbook()
    ws = wb.active
    ws.title = "Teacher portals"

    headers = ["#", "Teacher name (AR)", "Teacher name (Latin)", "Login ID", "Password", "Phone"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="F26522")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(creds, 1):
        ws.append(
            [
                i,
                row.get("name") or "",
                row.get("nameLatin") or "",
                row.get("loginCode") or "",
                row.get("password") or "",
                row.get("phone") or "",
            ]
        )

    widths = [6, 28, 28, 12, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
        for cell in row:
            cell.border = thin
            if cell.column in (1, 4, 5):
                cell.alignment = Alignment(horizontal="center")

    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.freeze_panes = "A2"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)

    # smoke: login first teacher with new credentials
    if creds:
        sample = creds[0]
        smoke = req(
            "/auth/login",
            {"loginCode": sample["loginCode"], "password": sample["password"]},
        )
        assert smoke.get("role") == "teacher", smoke
        print(
            "smoke login OK:",
            sample["loginCode"],
            sample["name"],
            "-> teacherId",
            smoke.get("teacherId"),
        )

    # also dump json backup
    backup = OUT.with_suffix(".json")
    backup.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", backup)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("FAILED:", e, file=sys.stderr)
        sys.exit(1)
